#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Otto.de 评论爬取工具（标题 + 正文合并分析版）
核心思路：
  · 把评论标题和正文拼成一段完整文本，一起检测标签、一起抽取关键句、一起生成总结
  · 标题里的问题（Fehlkauf / Nie wieder / Enttäuscht 等）会被主词库直接捕捉
  · 不再维护两套词库，逻辑简单、结果一致
"""

import os
import re
import time
import threading
from collections import Counter, defaultdict
from datetime import datetime
from html import unescape
from urllib.parse import urlparse

import requests
import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


# ============================================================
#  一、爬虫
# ============================================================
class OttoReviewScraper:
    def __init__(self):
        self.session = requests.Session()
        self.headers = {
            'User-Agent': ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
                           '(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36'),
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'de-DE,de;q=0.9,en-US;q=0.8,en;q=0.7',
            'Connection': 'keep-alive',
        }

    @staticmethod
    def extract_product_id(url_or_id):
        url_or_id = url_or_id.strip()
        if re.match(r'^[A-Z0-9]+$', url_or_id):
            return url_or_id
        parsed = urlparse(url_or_id)
        path = parsed.path.strip('/')
        m = re.search(r'kundenbewertungen/([A-Z0-9]+)', path, re.IGNORECASE)
        if m:
            return m.group(1)
        for part in reversed(path.split('/')):
            if re.match(r'^[A-Z0-9]+$', part):
                return part
        m = re.search(r'[?&]productId=([A-Z0-9]+)', url_or_id, re.IGNORECASE)
        if m:
            return m.group(1)
        raise ValueError(f"无法从链接中提取 product_id: {url_or_id}")

    def fetch_page(self, url):
        try:
            resp = self.session.get(url, headers=self.headers, timeout=30)
            resp.raise_for_status()
            resp.encoding = 'utf-8'
            return resp.text
        except requests.RequestException as e:
            print(f"请求失败: {e}")
            return None

    def parse_reviews(self, html, product_id):
        if not html:
            return []
        reviews = []
        item_pattern = r'<div\s+class="js_pdp_cr-item pdp_cr-item-content"[^>]*>'
        item_matches = list(re.finditer(item_pattern, html))
        if not item_matches:
            return []

        for i, match in enumerate(item_matches):
            start_pos = match.start()
            end_pos = item_matches[i + 1].start() if i + 1 < len(item_matches) else len(html)
            block = html[start_pos:end_pos]
            start_tag = match.group(0)

            review = {
                'product_id': product_id,
                'scraped_at': datetime.now().isoformat(),
                'source_url': f"https://www.otto.de/kundenbewertungen/{product_id}/",
            }

            rm = re.search(r'data-rating="(\d)"', start_tag)
            if rm:
                review['rating'] = int(rm.group(1))
            else:
                filled = len(re.findall(r'type="rating-filled"', block))
                review['rating'] = filled if filled > 0 else None

            tm = re.search(r'<h3[^>]*>(.*?)</h3>', block)
            if not tm:
                continue
            title = unescape(tm.group(1))
            review['title'] = re.sub(r'<[^>]+>', '', title).strip()

            hp = (r'(\d+)(?:\s*<[^>]+>\s*)*\s+von\s+(?:\s*<[^>]+>\s*)*(\d+)'
                  r'(?:\s*<[^>]+>\s*)*\s+finden\s+diese\s+Bewertung\s+hilfreich')
            hm = re.search(hp, block)
            if hm:
                review['helpful_count'] = int(hm.group(1))
                review['helpful_total'] = int(hm.group(2))
            else:
                review['helpful_count'] = None
                review['helpful_total'] = None

            review['verified_purchase'] = 'Verifizierter Kauf' in block

            txtm = re.search(r'pdp_cr-item__reviewText[^>]*>(.*?)</span>', block, re.DOTALL)
            if txtm:
                text = unescape(txtm.group(1))
                text = re.sub(r'<[^>]+>', ' ', text)
                review['review_text'] = ' '.join(text.split())
            else:
                review['review_text'] = ''

            author_patterns = [
                r'von\s+<b>([^<]+)</b>\s+aus\s+([^<]+?)\s+am\s+([\d\.]+)',
                r'(?:eine|von)\s+<b>([^<]+)</b>\s+aus\s+([^<]+?)\s+am\s+([\d\.]+)',
                r'eine\s+Kundenbewertung\s+aus\s+([^<]+?)\s+am\s+([\d\.]+)',
                r'von\s+([^<]+?)\s+aus\s+([^<]+?)\s+am\s+([\d\.]+)',
            ]
            review['author'] = review['location'] = review['date'] = None
            for pattern in author_patterns:
                m = re.search(pattern, block)
                if m:
                    if 'Kundenbewertung' in pattern:
                        review['author'] = 'Anonym'
                        review['location'] = m.group(1).strip()
                        review['date'] = m.group(2).strip()
                    else:
                        review['author'] = m.group(1).strip()
                        review['location'] = m.group(2).strip()
                        review['date'] = m.group(3).strip()
                    break

            variant = {}
            for key, label in [('size', 'Größe'), ('color', 'Farbe'),
                               ('variant', 'Ausführung'), ('material', 'Material')]:
                vm = re.search(rf'{label}:\s*<b>([^<]+)</b>', block)
                if vm:
                    variant[key] = unescape(vm.group(1).strip())
            review['product_variant'] = variant if variant else None

            sm = re.search(r'Verkäufer:\s*<b>([^<]+)</b>', block)
            review['seller'] = unescape(sm.group(1).strip()) if sm else None

            reviews.append(review)

        return reviews

    def get_all_reviews(self, product_id, max_pages=10, progress_callback=None):
        base_url = f"https://www.otto.de/kundenbewertungen/{product_id}/"
        all_reviews = []
        seen = set()

        for page in range(1, max_pages + 1):
            url = base_url if page == 1 else f"{base_url}?page={page}"
            if progress_callback:
                progress_callback(f"获取第 {page} 页 ...")

            html = self.fetch_page(url)
            if not html:
                break

            reviews = self.parse_reviews(html, product_id)
            if not reviews:
                if progress_callback:
                    progress_callback(f"第 {page} 页无评论，结束")
                break

            new_count = 0
            for r in reviews:
                key = (
                    r.get('title', ''),
                    r.get('author', ''),
                    r.get('date', ''),
                    (r.get('review_text', '') or '')[:60],
                )
                if key in seen:
                    continue
                seen.add(key)
                all_reviews.append(r)
                new_count += 1

            if new_count == 0:
                if progress_callback:
                    progress_callback("无新增评论，结束")
                break

            if progress_callback:
                progress_callback(f"第 {page} 页新增 {new_count} 条，累计 {len(all_reviews)} 条")

            time.sleep(1.2)

        return all_reviews


# ============================================================
#  二、标签库（合并版：标题与正文共用一套）
# ============================================================
ISSUE_PATTERNS = {
    "质量缺陷/破损": [
        r'defekt', r'kaputt', r'gebrochen', r'beschädigt', r'gerissen',
        r'eingerissen', r'aufgeplatzt', r'mangelhaft', r'fehlerhaft',
        r'schlechte qualität', r'keine qualität', r'nicht haltbar',
        r'schlecht verarbeitet', r'qualität mangelhaft', r'minderwertig',
    ],
    "稳定性差/摇晃": [
        r'wackel', r'kippel', r'kippt', r'instabil', r'schwankt',
        r'schief', r'verbiegt', r'nicht stabil', r'steht nicht',
    ],
    "异味/刺鼻": [
        r'geruch', r'riecht', r'gestank', r'stinkt', r'chemisch',
        r'unangenehmer geruch',
    ],
    "尺寸不符/偏小": [
        r'zu klein', r'zu eng', r'zu groß', r'passt nicht', r'zu kurz',
        r'zu schmal', r'zu niedrig', r'zu hoch', r'falsche größe',
        r'größe stimmt nicht', r'nicht passend', r'zu weit',
    ],
    "色差/与图不符": [
        r'farbabweichung', r'farbton', r'andere farbe',
        r'anders als abgebildet', r'anders als auf',
        r'falsche farbe', r'farbe stimmt nicht',
    ],
    "材质廉价/单薄": [
        r'billig', r'zu dünn', r'plastik', r'kunststoff',
        r'fühlt sich billig', r'billig verarbeitet',
    ],
    "组装/安装困难": [
        r'montage', r'aufbau', r'anleitung', r'aufbauen', r'kompliziert',
        r'bohrung', r'passt nicht zusammen', r'schwer aufzubauen',
        r'schlechte anleitung',
    ],
    "舒适度差": [
        r'unbequem', r'zu hart', r'druckstellen', r'rückenschmerz', r'tut weh',
        r'nicht bequem',
    ],
    "与描述不符": [
        r'nicht wie beschrieben', r'anders als beschrieben', r'nicht so wie',
        r'täuschend', r'irreführend', r'entspricht nicht der beschreibung',
    ],
    "客服/物流问题": [
        r'service', r'reklamation', r'lieferung', r'versand', r'hermes',
        r'dhl', r'zu spät', r'nicht geliefert', r'späte lieferung',
        r'nie angekommen', r'schlechter service',
    ],
    "性价比低": [
        r'überteuert', r'nicht wert', r'geld nicht wert', r'zu teuer',
        r'preis nicht gerechtfertigt',
    ],
    # ---------- 差评标题里常见、直接表达不满的泛负面表达 ----------
    "差评/不推荐": [
        r'nie wieder', r'finger weg', r'fehlkauf', r'enttäusch',
        r'katastrophe', r'desaster', r'unzufrieden', r'müll', r'schrott',
        r'nicht empfehlenswert', r'keine kaufempfehlung',
        r'nicht zu empfehlen', r'nicht zufrieden',
        r'sehr schlecht', r'ganz schlecht', r'echt schlecht',
        r'umtausch', r'retoure', r'zurückgeschickt', r'abzocke', r'betrug',
    ],
}

HIGHLIGHT_PATTERNS = {
    "质量好评": [
        r'gute qualität', r'hochwertig', r'solide', r'robust', r'stabil',
        r'gute verarbeitung', r'wertig', r'top qualität',
    ],
    "外观好评": [
        r'schön', r'sieht gut aus', r'optik', r'design', r'edel',
        r'wie abgebildet', r'sehr schick',
    ],
    "舒适好评": [r'bequem', r'komfortabel', r'angenehm', r'weich'],
    "安装简便": [
        r'einfach aufzubauen', r'schnell aufgebaut', r'leicht montiert',
        r'anleitung gut', r'leicht zu montieren',
    ],
    "性价比高": [
        r'preis-leistung', r'günstig', r'preiswert', r'sein geld wert',
        r'gutes preis', r'preis\s*/\s*leistung',
    ],
    "物流好评": [r'schnelle lieferung', r'pünktlich', r'gut verpackt', r'schneller versand'],
    "颜色满意": [r'farbe wie', r'farbe schön', r'farbe stimmt', r'farbe sehr schön'],
    "尺寸合适": [r'größe passt', r'passt perfekt', r'wie erwartet', r'genau richtig'],
    "推荐购买": [
        r'empfehlenswert', r'weiterempfehlen', r'würde wieder kaufen',
        r'klare kaufempfehlung', r'kann ich empfehlen',
    ],
}

EMPTY_STARS = {i: 0 for i in range(1, 6)}


# ============================================================
#  三、标签检测（单一入口，接收合并后的完整文本）
# ============================================================
def detect_tags(text):
    """
    从一段文本中检测问题标签、亮点标签、命中片段
    text 应为「标题 + 正文」合并后的完整文本
    """
    if not text:
        return [], [], {}
    issues, highlights, snippets = [], [], {}

    for tag, patterns in ISSUE_PATTERNS.items():
        for p in patterns:
            m = re.search(p, text, re.IGNORECASE)
            if m:
                if tag not in issues:
                    issues.append(tag)
                    if tag not in snippets:
                        s = max(0, m.start() - 30)
                        e = min(len(text), m.end() + 60)
                        snippets[tag] = text[s:e].strip()
                break

    for tag, patterns in HIGHLIGHT_PATTERNS.items():
        for p in patterns:
            m = re.search(p, text, re.IGNORECASE)
            if m:
                if tag not in highlights:
                    highlights.append(tag)
                    if tag not in snippets:
                        s = max(0, m.start() - 30)
                        e = min(len(text), m.end() + 60)
                        snippets[tag] = text[s:e].strip()
                break

    return issues, highlights, snippets


# ============================================================
#  四、内容总结生成
# ============================================================
def extract_key_sentence(text, tags):
    """从合并文本中挑出与命中标签最相关的一句原文"""
    if not text:
        return ""
    sentences = re.split(r'(?<=[.!?])\s+', text.strip())
    sentences = [s.strip() for s in sentences if len(s.strip()) > 3]
    if not sentences:
        return text.strip()[:150]

    all_patterns = []
    for tag in tags:
        if tag in ISSUE_PATTERNS:
            all_patterns.extend(ISSUE_PATTERNS[tag])
        if tag in HIGHLIGHT_PATTERNS:
            all_patterns.extend(HIGHLIGHT_PATTERNS[tag])

    if not all_patterns:
        return sentences[0][:150]

    best, best_score = "", 0
    for s in sentences:
        sl = s.lower()
        score = 0
        for p in all_patterns:
            if re.search(p, sl, re.IGNORECASE):
                score += 1
        if score > best_score:
            best_score = score
            best = s

    return (best or sentences[0])[:150]


def summarize_review(title, text, rating, issues, highlights):
    """
    对「标题 + 正文」合并后的内容生成总结
    格式：星级定性 ｜ 标题：「原标题」 ｜ 称赞：... ｜ 问题：... ｜ 原文：「关键句」
    标签来自标题+正文合并检测，不再区分来源
    """
    title = (title or '').strip()
    text = (text or '').strip()

    # 合并文本（标签检测与关键句抽取都基于它）
    if title and text:
        combined = f"{title}. {text}"
    else:
        combined = title or text
    if not combined:
        return ""

    parts = []

    # 1) 星级定性
    if isinstance(rating, int):
        if rating >= 4:
            parts.append(f"{rating}星好评")
        elif rating == 3:
            parts.append("3星中评")
        else:
            parts.append(f"{rating}星差评")
    else:
        parts.append("未评分")

    # 2) 标题原文（作为买家自述的核心观点保留）
    if title:
        parts.append(f"标题：「{title}」")

    # 3) 命中亮点（来自合并文本）
    if highlights:
        parts.append("称赞：" + "、".join(highlights))

    # 4) 命中问题（来自合并文本）
    if issues:
        parts.append("问题：" + "、".join(issues))

    # 5) 关键原文句（从合并文本里抽）
    key = extract_key_sentence(combined, (issues or []) + (highlights or []))
    if key and len(key) > 3:
        parts.append(f"原文：「{key}」")

    # 6) 完全没有标签时，直接引用原文
    if len(parts) <= 2:
        parts.append(f"原文：「{combined[:120]}」")

    return " ｜ ".join(parts)


# ============================================================
#  五、产品分析（标题+正文合并检测）
# ============================================================
def analyze_product(reviews):
    total = len(reviews)
    star_dist = dict(EMPTY_STARS)
    rating_sum = 0
    rated = 0

    for r in reviews:
        title = (r.get('title') or '').strip()
        text = (r.get('review_text') or '').strip()

        # 合并标题 + 正文为一段完整文本
        if title and text:
            combined = f"{title}. {text}"
        else:
            combined = title or text

        # 一次性检测标签
        issues, highlights, snippets = detect_tags(combined)
        r['issue_tags'] = issues
        r['highlight_tags'] = highlights
        r['tag_snippets'] = snippets
        r['combined_text'] = combined
        r['summary'] = summarize_review(title, text, r.get('rating'), issues, highlights)

        rt = r.get('rating')
        if isinstance(rt, int) and 1 <= rt <= 5:
            star_dist[rt] += 1
            rating_sum += rt
            rated += 1

    avg_rating = round(rating_sum / rated, 2) if rated else 0.0
    low_count = star_dist[1] + star_dist[2]
    low_ratio = low_count / rated if rated else 0.0

    # 变体维度（颜色 / 尺寸）
    variant_breakdown = {'color': [], 'size': []}
    for vkey in ('color', 'size'):
        bucket = defaultdict(lambda: {
            'total': 0, 'stars': dict(EMPTY_STARS), 'low': 0, 'samples': []
        })
        for r in reviews:
            v = r.get('product_variant') or {}
            val = v.get(vkey)
            rt = r.get('rating')
            if not val or not isinstance(rt, int) or not (1 <= rt <= 5):
                continue
            b = bucket[val]
            b['total'] += 1
            b['stars'][rt] += 1
            if rt <= 2:
                b['low'] += 1
                if len(b['samples']) < 2:
                    snip = r.get('summary', '') or ' '.join(
                        filter(None, [r.get('title') or '', r.get('review_text') or '']))
                    b['samples'].append(snip[:150])

        rows = []
        for val, b in bucket.items():
            rows.append({
                'value': val,
                'total': b['total'],
                'stars': b['stars'],
                'low': b['low'],
                'low_ratio': b['low'] / b['total'] if b['total'] else 0,
                'samples': b['samples'],
            })
        rows.sort(key=lambda x: (-x['low'], -x['low_ratio']))
        variant_breakdown[vkey] = rows

    worst_color = variant_breakdown['color'][0] if variant_breakdown['color'] else None
    worst_size = variant_breakdown['size'][0] if variant_breakdown['size'] else None
    if worst_color and worst_color['low'] == 0:
        worst_color = None
    if worst_size and worst_size['low'] == 0:
        worst_size = None

    # 低星原因（基于标签统计，标签来自标题+正文合并）
    reason_stats = defaultdict(lambda: {'count': 0, 'snippets': []})
    unmatched = 0
    for r in reviews:
        rt = r.get('rating')
        if not (isinstance(rt, int) and rt <= 2):
            continue
        tags = r.get('issue_tags', [])
        if not tags:
            unmatched += 1
        for tag in tags:
            reason_stats[tag]['count'] += 1
            snip = r.get('tag_snippets', {}).get(tag, '')
            if snip and len(reason_stats[tag]['snippets']) < 3:
                reason_stats[tag]['snippets'].append(snip)

    reasons = [{'issue': k, 'count': v['count'], 'snippets': v['snippets']}
               for k, v in reason_stats.items()]
    if unmatched:
        reasons.append({'issue': '其他/未归类', 'count': unmatched, 'snippets': []})
    reasons.sort(key=lambda x: x['count'], reverse=True)

    return {
        'total': total,
        'rated': rated,
        'star_dist': star_dist,
        'avg_rating': avg_rating,
        'low_count': low_count,
        'low_ratio': low_ratio,
        'variants': variant_breakdown,
        'worst_color': worst_color,
        'worst_size': worst_size,
        'reasons': reasons,
        'reviews': reviews,
    }


# ============================================================
#  六、Excel 导出
# ============================================================
HEADER_FILL = PatternFill("solid", start_color="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)
LOW_FILL = PatternFill("solid", start_color="FFE0E0")    # 1-2 星浅红
HIGH_FILL = PatternFill("solid", start_color="E3FBE3")   # 4-5 星浅绿
SUMMARY_FONT = Font(bold=False, color="1F4E79")


def _write_header(ws, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
    ws.freeze_panes = "A2"


def _autofit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_excel(analysis_by_product, output_dir, filename=None):
    if Workbook is None:
        return None
    if not filename:
        filename = os.path.join(
            output_dir,
            f"otto_review_analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )

    wb = Workbook()

    # ---------- Sheet1: 评论明细+总结（核心表） ----------
    ws1 = wb.active
    ws1.title = "评论明细+总结"
    headers1 = ["产品ID", "星级", "标题", "评论内容", "内容总结",
                "颜色", "尺寸", "规格", "材质",
                "作者", "日期", "验证购买",
                "问题标签", "亮点标签"]
    _write_header(ws1, headers1)

    for pid, a in analysis_by_product.items():
        # 组内按星级升序（低星在前），未知星级放最后
        sorted_reviews = sorted(
            a['reviews'],
            key=lambda r: (r.get('rating') if isinstance(r.get('rating'), int) else 99)
        )
        for r in sorted_reviews:
            v = r.get('product_variant') or {}
            ws1.append([
                r.get('product_id', ''),
                r.get('rating', ''),
                r.get('title', ''),
                r.get('review_text', ''),
                r.get('summary', ''),
                v.get('color', ''),
                v.get('size', ''),
                v.get('variant', ''),
                v.get('material', ''),
                r.get('author', ''),
                r.get('date', ''),
                '是' if r.get('verified_purchase') else '否',
                '、'.join(r.get('issue_tags', [])),
                '、'.join(r.get('highlight_tags', [])),
            ])

            rt = r.get('rating')
            if isinstance(rt, int):
                fill = LOW_FILL if rt <= 2 else (HIGH_FILL if rt >= 4 else None)
                if fill:
                    row_idx = ws1.max_row
                    for col in range(1, len(headers1) + 1):
                        ws1.cell(row=row_idx, column=col).fill = fill

    for row in ws1.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            if c.column == 5:
                c.font = SUMMARY_FONT
    _autofit(ws1, [14, 6, 26, 50, 80, 14, 12, 14, 14, 14, 12, 9, 26, 24])
    ws1.auto_filter.ref = ws1.dimensions

    # ---------- Sheet2: 商品总览 ----------
    ws2 = wb.create_sheet("商品总览")
    headers2 = ["产品ID", "评论总数", "平均星级", "5星", "4星", "3星", "2星", "1星",
                "差评数(1-2星)", "差评率", "差评最多的颜色", "差评最多的尺寸", "主要低星原因"]
    _write_header(ws2, headers2)

    for pid, a in analysis_by_product.items():
        wc = a['worst_color']
        ws_ = a['worst_size']
        top_reasons = '、'.join(f"{r['issue']}({r['count']})" for r in a['reasons'][:5])
        ws2.append([
            pid,
            a['total'],
            a['avg_rating'],
            a['star_dist'][5], a['star_dist'][4], a['star_dist'][3],
            a['star_dist'][2], a['star_dist'][1],
            a['low_count'],
            f"{a['low_ratio'] * 100:.1f}%",
            f"{wc['value']}（{wc['low']}/{wc['total']}）" if wc else "-",
            f"{ws_['value']}（{ws_['low']}/{ws_['total']}）" if ws_ else "-",
            top_reasons or "-",
        ])
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    _autofit(ws2, [14, 10, 10, 7, 7, 7, 7, 7, 14, 10, 28, 28, 48])
    ws2.auto_filter.ref = ws2.dimensions

    # ---------- Sheet3: 变体星级分布 ----------
    ws3 = wb.create_sheet("变体星级分布")
    headers3 = ["产品ID", "变体类型", "变体值", "评论数", "平均星级",
                "5星", "4星", "3星", "2星", "1星", "差评数(1-2星)", "差评率",
                "低星示例总结"]
    _write_header(ws3, headers3)

    for pid, a in analysis_by_product.items():
        for vkey, vlabel in (('color', '颜色'), ('size', '尺寸')):
            for row in a['variants'][vkey]:
                st = row['stars']
                rated = sum(st.values())
                avg = round(sum(k * v for k, v in st.items()) / rated, 2) if rated else 0
                samples = ' ｜ '.join(row.get('samples', [])[:2])
                ws3.append([
                    pid, vlabel, row['value'], row['total'], avg,
                    st[5], st[4], st[3], st[2], st[1],
                    row['low'], f"{row['low_ratio'] * 100:.1f}%",
                    samples,
                ])
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    _autofit(ws3, [14, 10, 22, 9, 10, 7, 7, 7, 7, 7, 14, 9, 72])
    ws3.auto_filter.ref = ws3.dimensions

    # ---------- Sheet4: 低星原因归纳 ----------
    ws4 = wb.create_sheet("低星原因归纳")
    headers4 = ["产品ID", "问题类别", "提及次数", "示例评论片段"]
    _write_header(ws4, headers4)

    for pid, a in analysis_by_product.items():
        if not a['reasons']:
            ws4.append([pid, "（无低星评论）", 0, ""])
            continue
        for item in a['reasons']:
            ws4.append([
                pid, item['issue'], item['count'],
                ' ｜ '.join(item['snippets'][:2])
            ])
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    _autofit(ws4, [14, 22, 12, 88])
    ws4.auto_filter.ref = ws4.dimensions

    wb.save(filename)
    return filename


# ============================================================
#  七、GUI
# ============================================================
class OttoScraperGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("Otto.de 评论爬取工具（标题+正文合并分析）")
        self.root.geometry("840x710")
        self.output_dir = os.getcwd()
        self.all_reviews = []

        style = ttk.Style()
        style.configure('TLabel', font=('微软雅黑', 11))
        style.configure('TButton', font=('微软雅黑', 11))

        main = ttk.Frame(root, padding="18")
        main.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        ttk.Label(main, text="Otto 产品链接或 Product ID（每行一个）:",
                  font=('微软雅黑', 12, 'bold')).grid(row=0, column=0, columnspan=2,
                                                     sticky=tk.W, pady=(0, 5))

        self.input_text = scrolledtext.ScrolledText(main, width=78, height=6,
                                                    font=('Consolas', 10))
        self.input_text.grid(row=1, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=(0, 4))
        self.input_text.insert(tk.END, "https://www.otto.de/kundenbewertungen/S0MDG0GH/")

        ttk.Label(main,
                  text="示例: S0MDG0GH  或  https://www.otto.de/kundenbewertungen/S0MDG0GH/",
                  font=('微软雅黑', 9), foreground='gray').grid(
            row=2, column=0, columnspan=2, sticky=tk.W, pady=(0, 12))

        sf = ttk.LabelFrame(main, text="设置", padding="10")
        sf.grid(row=3, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=(0, 12))

        ttk.Label(sf, text="每个产品最大页数:").grid(row=0, column=0, sticky=tk.W)
        self.pages_var = tk.StringVar(value="10")
        ttk.Spinbox(sf, from_=1, to=100, textvariable=self.pages_var,
                    width=8).grid(row=0, column=1, sticky=tk.W, padx=(6, 25))

        ttk.Label(sf, text="保存位置:").grid(row=0, column=2, sticky=tk.W)
        self.dir_label = ttk.Label(sf, text=self.output_dir, font=('微软雅黑', 9))
        self.dir_label.grid(row=0, column=3, sticky=tk.W, padx=(6, 5))
        ttk.Button(sf, text="更改目录", command=self.choose_directory).grid(
            row=0, column=4, sticky=tk.W)

        bf = ttk.Frame(main)
        bf.grid(row=4, column=0, columnspan=2, pady=(0, 10))
        self.start_btn = ttk.Button(bf, text="▶ 开始批量爬取并分析",
                                    command=self.start_scrape, width=24)
        self.start_btn.grid(row=0, column=0, padx=(0, 10))
        ttk.Button(bf, text="🗑 清空日志", command=self.clear_log, width=15).grid(row=0, column=1)

        self.progress = ttk.Progressbar(main, mode='indeterminate', length=780)
        self.progress.grid(row=5, column=0, columnspan=2, sticky=(tk.W, tk.E), pady=(0, 10))

        ttk.Label(main, text="运行日志:", font=('微软雅黑', 11, 'bold')).grid(
            row=6, column=0, sticky=tk.W, pady=(0, 4))

        self.log_text = scrolledtext.ScrolledText(main, width=90, height=18,
                                                  font=('Consolas', 10),
                                                  wrap=tk.WORD, state=tk.DISABLED)
        self.log_text.grid(row=7, column=0, columnspan=2, sticky=(tk.W, tk.E, tk.N, tk.S))

        self.status_var = tk.StringVar(value="就绪")
        ttk.Label(root, textvariable=self.status_var, relief=tk.SUNKEN,
                  anchor=tk.W, font=('微软雅黑', 9)).grid(row=1, column=0, sticky=(tk.W, tk.E))

        root.columnconfigure(0, weight=1)
        root.rowconfigure(0, weight=1)
        main.columnconfigure(0, weight=1)
        main.columnconfigure(1, weight=1)

        self.log("Otto.de 评论爬取工具已启动")
        self.log("请粘贴多个链接或 Product ID（每行一个），点击【开始批量爬取并分析】")
        self.log("标题与正文合并成完整文本一起分析，标题里的问题不会被遗漏")

    def log(self, message):
        self.log_text.configure(state=tk.NORMAL)
        self.log_text.insert(tk.END, f"[{datetime.now().strftime('%H:%M:%S')}] {message}\n")
        self.log_text.see(tk.END)
        self.log_text.configure(state=tk.DISABLED)

    def ui_log(self, message):
        self.root.after(0, lambda m=message: self.log(m))

    def clear_log(self):
        self.log_text.configure(state=tk.NORMAL)
        self.log_text.delete(1.0, tk.END)
        self.log_text.configure(state=tk.DISABLED)

    def choose_directory(self):
        d = filedialog.askdirectory(initialdir=self.output_dir)
        if d:
            self.output_dir = d
            self.dir_label.config(text=d)
            self.log(f"输出目录已更改为: {d}")

    def start_scrape(self):
        raw = self.input_text.get("1.0", tk.END).strip()
        if not raw:
            messagebox.showwarning("提示", "请输入至少一个产品链接或 Product ID！")
            return
        entries = [ln.strip() for ln in raw.splitlines() if ln.strip()]
        if not entries:
            messagebox.showwarning("提示", "请输入至少一个产品链接或 Product ID！")
            return
        try:
            max_pages = int(self.pages_var.get())
        except ValueError:
            messagebox.showwarning("提示", "页数必须是数字！")
            return

        self.start_btn.config(state=tk.DISABLED)
        self.progress.start()
        self.status_var.set("正在爬取...")
        self.clear_log()
        self.all_reviews = []

        t = threading.Thread(target=self._worker, args=(entries, max_pages), daemon=True)
        t.start()

    def _worker(self, entries, max_pages):
        by_product = defaultdict(list)
        errors = []

        self.ui_log(f"🔄 开始批量爬取，共 {len(entries)} 个产品")

        for idx, entry in enumerate(entries, 1):
            self.ui_log(f"\n--- [{idx}/{len(entries)}] {entry} ---")
            try:
                scraper = OttoReviewScraper()
                pid = scraper.extract_product_id(entry)
                self.ui_log(f"  product_id: {pid}")

                def cb(msg):
                    self.ui_log(f"    {msg}")

                reviews = scraper.get_all_reviews(pid, max_pages=max_pages,
                                                  progress_callback=cb)
                if not reviews:
                    self.ui_log("  ⚠️ 未获取到评论，跳过")
                    errors.append((entry, "无评论"))
                    continue

                by_product[pid].extend(reviews)
                self.ui_log(f"  ✅ 获取 {len(reviews)} 条评论")

            except Exception as e:
                self.ui_log(f"  ❌ 处理失败: {e}")
                errors.append((entry, str(e)))

        if not by_product:
            self.ui_log("\n❌ 没有任何数据，无法生成报告")
            self.root.after(0, self._finished)
            return

        self.ui_log("\n" + "=" * 55)
        self.ui_log("📊 开始分析（标题 + 正文合并检测）...")

        analysis_by_product = {}
        for pid, reviews in by_product.items():
            a = analyze_product(reviews)
            analysis_by_product[pid] = a

            st = a['star_dist']
            self.ui_log(f"\n【{pid}】共 {a['total']} 条评论，平均 {a['avg_rating']} 星")
            self.ui_log(f"  星级分布： 5★={st[5]}  4★={st[4]}  3★={st[3]}  "
                        f"2★={st[2]}  1★={st[1]}")
            self.ui_log(f"  差评(1-2星)： {a['low_count']} 条，占 {a['low_ratio'] * 100:.1f}%")

            wc = a['worst_color']
            if wc:
                self.ui_log(f"  差评最集中的颜色： {wc['value']} "
                            f"（{wc['low']}/{wc['total']} 条差评）")
            w = a['worst_size']
            if w:
                self.ui_log(f"  差评最集中的尺寸： {w['value']} "
                            f"（{w['low']}/{w['total']} 条差评）")

            if a['reasons']:
                top = '、'.join(f"{r['issue']}({r['count']})" for r in a['reasons'][:5])
                self.ui_log(f"  主要低星原因： {top}")
            else:
                self.ui_log("  暂无低星评论")

        self.ui_log("\n💾 正在生成 Excel 报告 ...")
        try:
            out_file = export_excel(analysis_by_product, self.output_dir)
            if out_file:
                self.ui_log(f"✅ 报告已生成： {out_file}")
                self.ui_log("   → 「评论明细+总结」表：标题与正文合并分析")
            else:
                self.ui_log("⚠️ 导出失败：请确认已安装 openpyxl（pip install openpyxl）")
        except Exception as e:
            out_file = None
            self.ui_log(f"❌ 导出 Excel 失败： {e}")

        if errors:
            self.ui_log("\n⚠️ 以下条目处理失败：")
            for e, msg in errors:
                self.ui_log(f"  {e} -> {msg}")

        total_reviews = sum(len(v) for v in by_product.values())
        self.ui_log(f"\n✅ 全部完成！成功产品 {len(by_product)} 个，"
                    f"总评论 {total_reviews} 条")
        self.ui_log(f"📁 输出目录： {self.output_dir}")

        self.root.after(0, lambda: messagebox.showinfo(
            "完成",
            f"批量爬取完成！\n\n"
            f"成功产品：{len(by_product)} 个\n"
            f"总评论数：{total_reviews} 条\n"
            f"报告位置：{out_file or self.output_dir}"
        ))
        self.root.after(0, self._finished)

    def _finished(self):
        self.progress.stop()
        self.start_btn.config(state=tk.NORMAL)
        self.status_var.set("就绪")


# ============================================================
#  八、入口
# ============================================================
def main():
    root = tk.Tk()
    OttoScraperGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()